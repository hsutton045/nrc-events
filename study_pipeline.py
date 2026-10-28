"""NRC/INL study discovery, content-addressed archive, auditable extraction staging.

No extracted numeric value is approved automatically. Run on a host with outbound HTTPS.
"""
import argparse, csv, hashlib, io, json, os, re, sys
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse
import requests
from bs4 import BeautifulSoup
from database import connect

SEEDS = [
 'https://nrcoe.inl.gov/avgperf/',
 'https://nrcoe.inl.gov/CompPerf',
 'https://nrcoe.inl.gov/SysStudy',
 'https://nrcoe.inl.gov/ReactorYears',
 'https://nrcoe.inl.gov/ccf_insights',
]
ROOT = Path(os.getenv('NRC_STUDY_ARCHIVE', str(Path(os.getenv('NRC_DB_PATH','data/nrc_events.db')).parent / 'studies')))
MAX_BYTES = 35 * 1024 * 1024
SCHEMA = '''
CREATE TABLE IF NOT EXISTS study_documents (
 url TEXT PRIMARY KEY, title TEXT NOT NULL, listing_url TEXT,
 sha256 TEXT, local_path TEXT, media_type TEXT, bytes INTEGER,
 checked_at TEXT, downloaded_at TEXT, extraction_status TEXT DEFAULT 'not_extracted',
 extraction_error TEXT
);
CREATE TABLE IF NOT EXISTS study_extraction_rows (
 id TEXT PRIMARY KEY, document_url TEXT NOT NULL REFERENCES study_documents(url),
 sheet_or_page TEXT NOT NULL, row_number INTEGER NOT NULL,
 raw_cells_json TEXT NOT NULL, candidate_component TEXT, candidate_failure_mode TEXT,
 candidate_value TEXT, candidate_unit TEXT,
 review_status TEXT NOT NULL DEFAULT 'needs_review', created_at TEXT DEFAULT CURRENT_TIMESTAMP
);
CREATE INDEX IF NOT EXISTS idx_study_extraction_document ON study_extraction_rows(document_url);
'''

def init():
 with connect() as c: c.executescript(SCHEMA)

def allowed(url):
 p=urlparse(url)
 return p.scheme=='https' and p.hostname in ('nrcoe.inl.gov','www.nrcoe.inl.gov') and not p.username and not p.password

def discover(session, listing):
 if not allowed(listing): raise ValueError('Source outside NRC/INL allowlist')
 r=session.get(listing,timeout=30); r.raise_for_status()
 soup=BeautifulSoup(r.text,'html.parser'); found=[]
 for a in soup.select('a[href]'):
  u=urljoin(listing,a['href']).split('#')[0]
  if allowed(u) and re.search(r'\.(pdf|xlsx?|csv)(?:\?|$)',u,re.I):
   found.append((u,a.get_text(' ',strip=True)[:240] or Path(urlparse(u).path).name,listing))
 return list(dict(((u,t,l),None) for u,t,l in found))

def register(links):
 init()
 with connect() as c:
  for url,title,listing in links:
   c.execute('''INSERT INTO study_documents(url,title,listing_url) VALUES(?,?,?)
    ON CONFLICT(url) DO UPDATE SET title=excluded.title,listing_url=excluded.listing_url''',(url,title,listing))

def download(session,url):
 if not allowed(url):raise ValueError('Source outside allowlist')
 with session.get(url,timeout=90,stream=True,allow_redirects=True) as r:
  r.raise_for_status()
  if not allowed(r.url):raise ValueError('Redirect outside allowlist')
  buf=io.BytesIO()
  for chunk in r.iter_content(65536):
   buf.write(chunk)
   if buf.tell()>MAX_BYTES:raise ValueError('File exceeds download limit')
  raw=buf.getvalue()
  ext=Path(urlparse(r.url).path).suffix.lower()
  if ext not in ('.pdf','.xls','.xlsx','.csv'):raise ValueError('Unsupported file extension')
  signature={'.pdf':raw.startswith(b'%PDF'),'.xlsx':raw.startswith(b'PK'),'.xls':raw.startswith(bytes.fromhex('D0CF11E0')),'.csv':True}[ext]
  if not signature:raise ValueError('File signature does not match extension')
  digest=hashlib.sha256(raw).hexdigest();ROOT.mkdir(parents=True,exist_ok=True)
  path=ROOT / (digest+ext)
  if not path.exists():path.write_bytes(raw)
  now=datetime.now(timezone.utc).isoformat()
  with connect() as c:
   c.execute('''UPDATE study_documents SET sha256=?,local_path=?,media_type=?,bytes=?,checked_at=?,downloaded_at=? WHERE url=?''',
    (digest,str(path),ext,len(raw),now,now,url))
  return str(path)

def extract(url,path):
 ext=Path(path).suffix.lower(); rows=[]
 if ext=='.xlsx':
  from openpyxl import load_workbook
  wb=load_workbook(path,read_only=True,data_only=True)
  for sheet in wb:
   for i,row in enumerate(sheet.iter_rows(values_only=True),1):
    cells=[str(v)[:400] if v is not None else '' for v in row[:40]]
    if any(cells):rows.append((sheet.title,i,cells))
 elif ext=='.csv':
  with open(path,encoding='utf-8-sig',errors='replace',newline='') as f:
   for i,row in enumerate(csv.reader(f),1):
    if any(row):rows.append(('CSV',i,row[:40]))
 elif ext=='.pdf':
  import pdfplumber
  with pdfplumber.open(path) as pdf:
   for page_no,page in enumerate(pdf.pages,1):
    for table_idx,table in enumerate(page.extract_tables() or [],1):
     for i,row in enumerate(table,1):
      cells=[str(v or '')[:400] for v in row[:40]]
      if any(cells):rows.append((f'page {page_no} table {table_idx}',i,cells))
 else:raise ValueError('Legacy .xls unsupported for automatic extraction; download preserved for manual conversion')
 with connect() as c:
  # Preserve reviewer decisions across repeat extraction of the same source version.
  c.execute("DELETE FROM study_extraction_rows WHERE document_url=? AND review_status='needs_review'",(url,))
  for section,n,cells in rows:
   rid=hashlib.sha256(json.dumps([url,section,n,cells],sort_keys=True).encode()).hexdigest()
   c.execute('''INSERT INTO study_extraction_rows(id,document_url,sheet_or_page,row_number,raw_cells_json)
    VALUES(?,?,?,?,?) ON CONFLICT(id) DO NOTHING''',(rid,url,section,n,json.dumps(cells)))
  c.execute("UPDATE study_documents SET extraction_status='staged',extraction_error=NULL WHERE url=?",(url,))
 return len(rows)

def sync(seeds=SEEDS,max_downloads=20):
 init();session=requests.Session();session.headers.update({'User-Agent':'NRC-Research-Study-Indexer/1.0 (academic research)'})
 out={'discovered':0,'downloaded':0,'unchanged':0,'staged_rows':0,'errors':[]}
 for seed in seeds:
  try:
   links=discover(session,seed);register(links);out['discovered']+=len(links)
  except Exception as e:out['errors'].append({'source':seed,'error':str(e)})
 with connect() as c:
  todo=[dict(x) for x in c.execute('SELECT * FROM study_documents ORDER BY CASE WHEN local_path IS NULL THEN 0 ELSE 1 END, url')]
 for doc in todo[:max_downloads]:
  try:
   path=download(session,doc['url']);digest=Path(path).stem
   if doc['sha256']==digest and doc['extraction_status']=='staged':out['unchanged']+=1;continue
   out['downloaded']+=1
   try:out['staged_rows']+=extract(doc['url'],path)
   except Exception as e:
    with connect() as c:c.execute('UPDATE study_documents SET extraction_status=?,extraction_error=? WHERE url=?',('failed',str(e)[:1000],doc['url']))
    out['errors'].append({'source':doc['url'],'extraction_error':str(e)})
  except Exception as e:out['errors'].append({'source':doc['url'],'error':str(e)})
 return out

def export_staging(path):
 init()
 with connect() as c:
  rows=c.execute('''SELECT d.title,d.url,r.sheet_or_page,r.row_number,r.raw_cells_json,r.review_status
  FROM study_extraction_rows r JOIN study_documents d ON d.url=r.document_url ORDER BY d.url,r.sheet_or_page,r.row_number''').fetchall()
 with open(path,'w',newline='',encoding='utf-8') as f:
  w=csv.writer(f);w.writerow(['study_title','source_url','sheet_or_page','row_number','raw_cells_json','review_status'])
  for r in rows:w.writerow(tuple(r))
 return len(rows)

if __name__=='__main__':
 p=argparse.ArgumentParser();sub=p.add_subparsers(dest='cmd',required=True)
 s=sub.add_parser('sync');s.add_argument('--max-downloads',type=int,default=20)
 e=sub.add_parser('export-staging');e.add_argument('csv_path')
 a=p.parse_args();print(json.dumps(sync(max_downloads=a.max_downloads) if a.cmd=='sync' else {'rows':export_staging(a.csv_path)},indent=2))
